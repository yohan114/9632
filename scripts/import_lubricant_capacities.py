import os
import sys
import sqlite3
import openpyxl

EXCEL_PATH = r"d:\Master system 1\Fleet_Oil_Lubricant_Capacities.xlsx"
DB_PATH = r"d:\Master system 1\data\workshopone.db"

def clean_val(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s != "" else None

def clean_num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None

def clean_int(v, default=0):
    if v is None or v == "":
        return default
    try:
        return int(float(v))
    except (ValueError, TypeError):
        return default

def main():
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: Excel file not found: {EXCEL_PATH}")
        sys.exit(1)

    print(f"Loading workbook: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Sheets found: {sheet_names}")

    print(f"Connecting to database: {DB_PATH}")
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()

    # Drop removed tables if they still linger
    cur.execute("DROP TABLE IF EXISTS lubricant_capacity_evidence")
    cur.execute("DROP TABLE IF EXISTS other_equipment_capacities")

    # Ensure vehicle_lubricant_capacities exists
    cur.execute("""
    CREATE TABLE IF NOT EXISTS vehicle_lubricant_capacities (
      id                   INTEGER PRIMARY KEY AUTOINCREMENT,
      asset_id             INTEGER REFERENCES assets(id) ON DELETE SET NULL,
      sheet_id             INTEGER,
      ec_no                TEXT,
      registration         TEXT,
      category             TEXT,
      brand                TEXT,
      model                TEXT,
      year                 TEXT,
      engine_oil_l         REAL,
      engine_oil_grade     TEXT,
      gearbox_oil_l        REAL,
      gearbox_oil_grade    TEXT,
      diff_oil_l           REAL,
      diff_oil_grade       TEXT,
      front_axle_oil_l     REAL,
      hydraulic_oil_l      REAL,
      final_drive_oil_l    REAL,
      swing_oil_l          REAL,
      other_gearbox_oil_l  REAL,
      coolant_l            REAL,
      brake_fluid_l        REAL,
      engine_oil_basis     TEXT,
      engine_oil_records   INTEGER DEFAULT 0,
      notes                TEXT,
      updated_by           TEXT,
      created_at           TEXT NOT NULL DEFAULT (datetime('now')),
      updated_at           TEXT NOT NULL DEFAULT (datetime('now'))
    )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_vlc_asset ON vehicle_lubricant_capacities(asset_id)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_vlc_ec ON vehicle_lubricant_capacities(ec_no)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_vlc_reg ON vehicle_lubricant_capacities(registration)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_vlc_cat ON vehicle_lubricant_capacities(category)")

    # Build asset lookup caches
    cur.execute("SELECT id, code, registration FROM assets")
    asset_by_code = {}
    asset_by_reg = {}
    for aid, code, reg in cur.fetchall():
        if code:
            asset_by_code[code.strip().upper()] = aid
        if reg:
            asset_by_reg[reg.strip().upper()] = aid

    # 1. Import Vehicle Capacities
    print("\n--- Importing Vehicle Capacities ---")
    ws_vc = wb['Vehicle Capacities']
    cur.execute("DELETE FROM vehicle_lubricant_capacities")

    vc_rows = 0
    for r in range(4, ws_vc.max_row + 1):
        sheet_id = clean_int(ws_vc.cell(r, 1).value, None)
        ec_no = clean_val(ws_vc.cell(r, 2).value)
        reg = clean_val(ws_vc.cell(r, 3).value)
        cat = clean_val(ws_vc.cell(r, 4).value)
        brand = clean_val(ws_vc.cell(r, 5).value)
        model = clean_val(ws_vc.cell(r, 6).value)
        year = clean_val(ws_vc.cell(r, 7).value)

        eng_l = clean_num(ws_vc.cell(r, 8).value)
        eng_grade = clean_val(ws_vc.cell(r, 9).value)
        gb_l = clean_num(ws_vc.cell(r, 10).value)
        gb_grade = clean_val(ws_vc.cell(r, 11).value)
        diff_l = clean_num(ws_vc.cell(r, 12).value)
        diff_grade = clean_val(ws_vc.cell(r, 13).value)
        faxle_l = clean_num(ws_vc.cell(r, 14).value)
        hyd_l = clean_num(ws_vc.cell(r, 15).value)
        fd_l = clean_num(ws_vc.cell(r, 16).value)
        swg_l = clean_num(ws_vc.cell(r, 17).value)
        ogb_l = clean_num(ws_vc.cell(r, 18).value)
        cool_l = clean_num(ws_vc.cell(r, 19).value)
        brk_l = clean_num(ws_vc.cell(r, 20).value)
        basis = clean_val(ws_vc.cell(r, 21).value)
        records = clean_int(ws_vc.cell(r, 22).value, 0)
        notes = clean_val(ws_vc.cell(r, 23).value)

        if not ec_no and not reg and not cat and not brand:
            continue

        aid = None
        if ec_no and ec_no.strip().upper() in asset_by_code:
            aid = asset_by_code[ec_no.strip().upper()]
        elif reg and reg.strip().upper() in asset_by_reg:
            aid = asset_by_reg[reg.strip().upper()]

        cur.execute("""
            INSERT INTO vehicle_lubricant_capacities (
              asset_id, sheet_id, ec_no, registration, category, brand, model, year,
              engine_oil_l, engine_oil_grade, gearbox_oil_l, gearbox_oil_grade,
              diff_oil_l, diff_oil_grade, front_axle_oil_l, hydraulic_oil_l,
              final_drive_oil_l, swing_oil_l, other_gearbox_oil_l, coolant_l, brake_fluid_l,
              engine_oil_basis, engine_oil_records, notes, updated_by
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            aid, sheet_id, ec_no, reg, cat, brand, model, year,
            eng_l, eng_grade, gb_l, gb_grade,
            diff_l, diff_grade, faxle_l, hyd_l,
            fd_l, swg_l, ogb_l, cool_l, brk_l,
            basis, records, notes, 'system_import'
        ))
        vc_rows += 1

    print(f"Imported {vc_rows} vehicle capacity records.")
    conn.commit()
    conn.close()
    print("\n[SUCCESS] VEHICLE CAPACITIES IMPORT COMPLETED SUCCESSFULLY!")

if __name__ == '__main__':
    main()
